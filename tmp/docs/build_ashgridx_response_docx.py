from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\PC\Documents\builtright")
OUTPUT = ROOT / "output" / "docx" / "BuiltRight_AshGridX_Pilot_Response.docx"
LOGO = ROOT / "src" / "assets" / "logoooo.png"

RED = "B32127"
TEAL = "159284"
INK = "17262B"
MUTED = "5D6B70"
PALE_TEAL = "EAF6F4"
PALE_RED = "FBEDEE"
LINE = "D9E4E5"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge in kwargs:
            edge_data = kwargs.get(edge)
            tag = "w:{}".format(edge)
            element = borders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                borders.append(element)
            for key in ["val", "sz", "space", "color"]:
                if key in edge_data:
                    element.set(qn("w:{}".format(key)), str(edge_data[key]))


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths_dxa, indent=120):
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")


def set_run(run, size=11, color=INK, bold=False, italic=False, font="Calibri"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:ascii"), font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def style_paragraph(paragraph, before=0, after=6, line=1.1, align=None):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if align is not None:
        paragraph.alignment = align


def add_text(doc, text, size=11, color=INK, bold=False, italic=False, before=0, after=6, line=1.1, align=None):
    paragraph = doc.add_paragraph()
    style_paragraph(paragraph, before, after, line, align)
    run = paragraph.add_run(text)
    set_run(run, size=size, color=color, bold=bold, italic=italic)
    return paragraph


def add_rich_paragraph(container, parts, before=0, after=6, line=1.1):
    paragraph = container.add_paragraph()
    style_paragraph(paragraph, before, after, line)
    for text, kwargs in parts:
        run = paragraph.add_run(text)
        set_run(run, **kwargs)
    return paragraph


def add_heading(doc, text, level=1):
    sizes = {1: 16, 2: 13, 3: 12}
    before = {1: 16, 2: 12, 3: 8}
    after = {1: 8, 2: 6, 3: 4}
    paragraph = doc.add_paragraph()
    style_paragraph(paragraph, before[level], after[level], 1.1)
    run = paragraph.add_run(text)
    set_run(run, size=sizes[level], color=TEAL if level < 3 else RED, bold=True)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    style_paragraph(paragraph, 0, 4, 1.167)
    paragraph.paragraph_format.left_indent = Inches(0.5)
    paragraph.paragraph_format.first_line_indent = Inches(-0.25)
    run = paragraph.add_run(text)
    set_run(run, size=11, color=INK)
    return paragraph


def add_numbered(doc, number, heading, detail):
    paragraph = doc.add_paragraph()
    style_paragraph(paragraph, 0, 0, 1.1)
    paragraph.paragraph_format.left_indent = Inches(0.55)
    paragraph.paragraph_format.first_line_indent = Inches(-0.55)
    marker = paragraph.add_run(f"{number:02d}  ")
    set_run(marker, size=11, color=RED, bold=True)
    title = paragraph.add_run(f"{heading}\n")
    set_run(title, size=11, color=INK, bold=True)
    detail_run = paragraph.add_run(detail)
    set_run(detail_run, size=10.5, color=MUTED)
    return paragraph


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run(run, size=9, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_card(table, row, col, title, body, fill, accent):
    cell = table.cell(row, col)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=130, bottom=130, start=140, end=140)
    set_cell_border(cell, left={"val": "single", "sz": 18, "color": accent}, top={"val": "single", "sz": 4, "color": LINE}, bottom={"val": "single", "sz": 4, "color": LINE}, right={"val": "single", "sz": 4, "color": LINE})
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    p1 = cell.paragraphs[0]
    style_paragraph(p1, 0, 3, 1.1)
    r1 = p1.add_run(title)
    set_run(r1, size=10.5, color=INK, bold=True)
    p2 = cell.add_paragraph()
    style_paragraph(p2, 0, 0, 1.1)
    r2 = p2.add_run(body)
    set_run(r2, size=9.5, color=INK)


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)

    header = section.header.paragraphs[0]
    style_paragraph(header, 0, 0, 1.0)
    left = header.add_run("BUILTRIGHT SERVICES LTD")
    set_run(left, size=8.5, color=TEAL, bold=True)
    right = header.add_run("                                      AshGridX pilot response")
    set_run(right, size=8.5, color=MUTED)

    footer = section.footer.paragraphs[0]
    style_paragraph(footer, 0, 0, 1.0)
    left = footer.add_run("BuiltRight Services Ltd | builtrightltd.com")
    set_run(left, size=8.5, color=MUTED)
    add_page_number(footer)

    # Masthead
    mast = doc.add_table(rows=1, cols=2)
    mast.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_width(mast, [2200, 7160], indent=0)
    mast.cell(0, 0).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    mast.cell(0, 1).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    mast.cell(0, 0).paragraphs[0].add_run().add_picture(str(LOGO), width=Inches(1.3))
    for cell in mast.rows[0].cells:
        set_cell_margins(cell, top=0, bottom=80, start=0, end=120)
        set_cell_border(cell, top={"val": "nil"}, bottom={"val": "nil"}, left={"val": "nil"}, right={"val": "nil"})
    p_title = mast.cell(0, 1).paragraphs[0]
    style_paragraph(p_title, 0, 2, 1.0)
    r = p_title.add_run("AshGridX Pilot Integration")
    set_run(r, size=23, color=INK, bold=True)
    p_sub = mast.cell(0, 1).add_paragraph()
    style_paragraph(p_sub, 0, 0, 1.0)
    r = p_sub.add_run("Technical clarification acknowledgement and next-step request")
    set_run(r, size=11, color=MUTED)

    # Metadata grid
    meta = doc.add_table(rows=2, cols=4)
    meta.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_width(meta, [1000, 3300, 1100, 3960], indent=120)
    metadata = [
        ("TO", "AshGridX Technical and Pilot Support Team", "DATE", "17 August 2026"),
        ("SUBJECT", "BuiltRight pilot integration - clarification acknowledgement and next steps", "FROM", "BuiltRight Services Ltd"),
    ]
    for i, row in enumerate(metadata):
        for j, text in enumerate(row):
            cell = meta.cell(i, j)
            set_cell_shading(cell, "F7FAFA")
            set_cell_margins(cell, top=90, bottom=70, start=110, end=110)
            set_cell_border(cell, top={"val": "single", "sz": 4, "color": LINE}, bottom={"val": "single", "sz": 4, "color": LINE}, left={"val": "single", "sz": 4, "color": LINE}, right={"val": "single", "sz": 4, "color": LINE})
            par = cell.paragraphs[0]
            style_paragraph(par, 0, 0, 1.0)
            run = par.add_run(text)
            set_run(run, size=9.5 if j % 2 == 0 else 10.5, color=MUTED if j % 2 == 0 else INK, bold=j % 2 == 0)

    add_text(doc, "Dear Israel,", after=7)
    add_text(doc, "Thank you for the detailed technical clarification. We have reviewed the responses and are ready to proceed with the BuiltRight pilot integration.", after=7)
    add_text(doc, "Our understanding of the confirmed AshGridX behavior is summarized below.", after=8)

    cards = doc.add_table(rows=2, cols=2)
    cards.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_width(cards, [4680, 4680], indent=120)
    add_card(cards, 0, 0, "Device identity", "BuiltRight will use deviceNumber as the primary identifier for every customer device. AshGridX will generate customerDeviceId during onboarding.", PALE_TEAL, TEAL)
    add_card(cards, 0, 1, "Protection and tamper", "A cable disconnection triggers a BYPASS/protection event and physically shuts down the inverter. Webhook delivery may take up to five minutes.", PALE_RED, RED)
    add_card(cards, 1, 0, "Control behavior", "A successful control response confirms that the request was accepted. BuiltRight will verify the actual state through the status endpoint. Offline commands are not queued.", PALE_TEAL, TEAL)
    add_card(cards, 1, 1, "Webhook security", "Webhooks will use HMAC-SHA256 verification with a 120-second timestamp tolerance. Protection removal is a separate operation requiring its own authorization key.", PALE_RED, RED)

    add_heading(doc, "Pilot activation requirements", 2)
    add_text(doc, "To begin the pilot, please provide the following items:", after=4)
    requirements = [
        "Pilot device purchase and onboarding instructions.",
        "Staging API base URL and API key.",
        "Webhook secret and the exact signature and timestamp header names.",
        "Required IP-whitelisting details.",
        "Device onboarding and assignment procedure.",
        "Complete staging API documentation, including status, control, protection-history, and device-management endpoints.",
        "Sample signed webhook payloads for BYPASS, restoration, online, and offline events.",
        "AshGridX pilot testing checklist and acceptance criteria.",
        "Confirmation of the support process during installation and testing.",
    ]
    for item in requirements:
        add_bullet(doc, item)

    doc.add_page_break()
    add_heading(doc, "BuiltRight implementation and test approach", 2)
    add_text(doc, "BuiltRight will map each customer to a unique deviceNumber, test tamper alerts, verify device state changes, and maintain a complete command and event audit trail.", after=7)
    add_text(doc, "Remote control will remain disabled until staging tests are successfully completed and approved by both parties.", after=10)

    steps = [
        (1, "Onboard", "Purchase and register the first pilot device through the AshGridX dashboard."),
        (2, "Connect", "Configure the staging credentials, webhook secret, callback URL, and any required IP whitelist."),
        (3, "Map", "Assign the AshGridX customerDeviceId to the BuiltRight deviceNumber and customer record."),
        (4, "Test", "Run signed BYPASS, restoration, online, offline, status, and control-state tests."),
        (5, "Approve", "Review the audit trail and jointly approve the pilot before enabling live control."),
    ]
    for number, heading, detail in steps:
        add_numbered(doc, number, heading, detail)

    callout = doc.add_table(rows=1, cols=1)
    callout.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_width(callout, [9360], indent=120)
    cell = callout.cell(0, 0)
    set_cell_shading(cell, PALE_RED)
    set_cell_margins(cell, top=130, bottom=130, start=150, end=150)
    set_cell_border(cell, top={"val": "single", "sz": 5, "color": "E6B9BC"}, bottom={"val": "single", "sz": 5, "color": "E6B9BC"}, left={"val": "single", "sz": 5, "color": "E6B9BC"}, right={"val": "single", "sz": 5, "color": "E6B9BC"})
    p = cell.paragraphs[0]
    style_paragraph(p, 0, 0, 1.0)
    r = p.add_run("Please let us know the next step for purchasing and onboarding the first pilot device.")
    set_run(r, size=11, color=RED, bold=True)

    add_text(doc, "Kind regards,", before=18, after=14)
    add_rich_paragraph(doc, [("BuiltRight Services Ltd\n", {"size": 11, "color": INK, "bold": True}), ("Operations and Technical Integration Team", {"size": 11, "color": INK})], after=12)
    add_text(doc, "BuiltRight Services Ltd | Lagos, Nigeria | builtrightltd.com", size=9, color=MUTED, after=0)

    doc.core_properties.title = "BuiltRight AshGridX Pilot Response"
    doc.core_properties.subject = "AshGridX pilot integration clarification acknowledgement and next steps"
    doc.core_properties.author = "BuiltRight Services Ltd"
    doc.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    build()
